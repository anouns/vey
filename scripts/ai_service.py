import sys
import os
import subprocess
import json
import time
import asyncio
from fastapi import FastAPI
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
import uvicorn
import psutil
import requests
from fastapi.middleware.cors import CORSMiddleware

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Global workspace state
CURRENT_WORKSPACE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SETTINGS_FILE = os.path.join(os.path.expanduser("~"), ".vey_settings.json")

# Boot timing
BOOT_START_TIME = time.time()
BOOT_STATUS = {
    "stage": "starting",
    "ready": False,
    "models_loaded": False,
    "workspace_ready": False,
    "uptime_sec": 0,
    "error": None,
}

# ─── LAZY MODEL LOADING ────────────────────────────────────────────
# Models are loaded on first use, not on startup, for instant server boot.
embed_model = None
llm_model = None
tokenizer = None
device = None
DOC_DB = []
EMBED_DB = None
LAST_INDEXED_PATH = None

def get_device():
    global device
    if device is None:
        import torch
        device = "cuda" if torch.cuda.is_available() else "cpu"
    return device

def get_embed_model():
    global embed_model
    if embed_model is None:
        from sentence_transformers import SentenceTransformer
        embed_model = SentenceTransformer('all-MiniLM-L6-v2', device=get_device())
    return embed_model

def get_llm():
    global llm_model, tokenizer
    if llm_model is None:
        import torch
        from transformers import AutoModelForCausalLM, AutoTokenizer
        model_id = "Qwen/Qwen2.5-Coder-1.5B-Instruct"
        tokenizer = AutoTokenizer.from_pretrained(model_id)
        d = get_device()
        llm_model = AutoModelForCausalLM.from_pretrained(
            model_id,
            torch_dtype=torch.float32 if d == "cpu" else torch.float16
        ).to(d)
    return llm_model, tokenizer

# ─── SETTINGS PERSISTENCE ──────────────────────────────────────────
def load_server_settings():
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            pass
    return {}

def save_server_settings(data):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

# On startup, restore last workspace
_saved = load_server_settings()
if _saved.get("last_workspace") and os.path.isdir(_saved["last_workspace"]):
    CURRENT_WORKSPACE = _saved["last_workspace"]

# Mark server as ready
BOOT_STATUS["stage"] = "ready"
BOOT_STATUS["ready"] = True
BOOT_STATUS["workspace_ready"] = True

# ─── VECTOR DATABASE ───────────────────────────────────────────────
def build_vector_database(root_dir):
    global DOC_DB, EMBED_DB, LAST_INDEXED_PATH
    import torch
    if root_dir == LAST_INDEXED_PATH and EMBED_DB is not None:
        return

    DOC_DB = []
    allowed_exts = {'.rs', '.py', '.js', '.jsx', '.ts', '.tsx', '.html', '.css',
                    '.toml', '.json', '.md', '.txt', '.pdf', '.docx', '.odt',
                    '.c', '.cpp', '.h', '.java', '.go', '.rb', '.php', '.sh', '.bat', '.ps1'}
    ignore_dirs = {'.git', 'node_modules', 'dist', 'build', '__pycache__', '.vey', 'target', '.next'}

    for root, dirs, filenames in os.walk(root_dir):
        dirs[:] = [d for d in dirs if d not in ignore_dirs]
        for f in filenames:
            ext = os.path.splitext(f)[1].lower()
            if ext in allowed_exts:
                file_path = os.path.join(root, f)
                try:
                    content = ""
                    if ext == '.pdf':
                        try:
                            import pypdf
                            with open(file_path, 'rb') as file:
                                reader = pypdf.PdfReader(file)
                                for page in reader.pages:
                                    t = page.extract_text()
                                    if t:
                                        content += t + "\n"
                        except:
                            pass
                    elif ext == '.docx':
                        try:
                            import docx
                            doc = docx.Document(file_path)
                            content = "\n".join([para.text for para in doc.paragraphs])
                        except:
                            pass
                    elif ext == '.odt':
                        try:
                            from odf import text as odf_text, teletype
                            from odf.opendocument import load
                            doc = load(file_path)
                            for para in doc.getElementsByType(odf_text.P):
                                content += teletype.extractText(para) + "\n"
                        except:
                            pass
                    else:
                        with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                            content = file.read()

                    if not content.strip():
                        continue

                    chunk_size = 900
                    rel_path = os.path.relpath(file_path, root_dir)
                    for i in range(0, len(content), chunk_size):
                        snippet = content[i:i+chunk_size]
                        chunk_text = f"--- ФАЙЛ: {rel_path} (ЧАСТЬ {i//chunk_size}) ---\n{snippet}\n"
                        DOC_DB.append(chunk_text)
                except:
                    continue

    if DOC_DB:
        model = get_embed_model()
        embeddings = model.encode(DOC_DB, convert_to_tensor=True)
        EMBED_DB = embeddings
    else:
        EMBED_DB = None
    LAST_INDEXED_PATH = root_dir


def retrieve_top_k(query, k=4):
    import torch
    if not DOC_DB or EMBED_DB is None:
        return "NO RELEVANT FILES."
    model = get_embed_model()
    query_emb = model.encode([query], convert_to_tensor=True)
    cos_scores = torch.nn.functional.cosine_similarity(query_emb, EMBED_DB)
    top_scores, top_idx = torch.topk(cos_scores, min(k, len(DOC_DB)))

    retrieved = []
    for idx in top_idx:
        retrieved.append(DOC_DB[idx.item()])
    return "\n".join(retrieved)


# ─── API MODELS ────────────────────────────────────────────────────
class ChatRequest(BaseModel):
    messages: list
    model: str = "PYTHON_LOCAL_AI"
    api_key: Optional[str] = None
    stream: bool = False

class TerminalRequest(BaseModel):
    command: str

class FileWriteRequest(BaseModel):
    filename: str
    content: str

class FileEditRequest(BaseModel):
    filepath: str
    content: str
    mode: str = "overwrite"  # "overwrite" | "append" | "patch"
    patch_target: Optional[str] = None
    patch_replacement: Optional[str] = None


# ─── BOOT STATUS ──────────────────────────────────────────────────
@app.get("/status")
async def get_status():
    """Return detailed boot/health status for the frontend boot screen."""
    return {
        **BOOT_STATUS,
        "uptime_sec": round(time.time() - BOOT_START_TIME, 1),
        "workspace": CURRENT_WORKSPACE,
        "pid": os.getpid(),
    }


# ─── FILE OPERATIONS ──────────────────────────────────────────────
@app.post("/file/write")
async def write_file(request: FileWriteRequest):
    try:
        global CURRENT_WORKSPACE
        # Support both relative and absolute paths
        if os.path.isabs(request.filename):
            target_path = request.filename
        else:
            target_path = os.path.join(CURRENT_WORKSPACE, request.filename)

        # Create parent directories if needed
        parent_dir = os.path.dirname(target_path)
        if parent_dir and not os.path.exists(parent_dir):
            os.makedirs(parent_dir, exist_ok=True)

        with open(target_path, "w", encoding="utf-8") as f:
            f.write(request.content)
        return {"status": "SUCCESS", "path": target_path}
    except Exception as e:
        return {"status": "ERROR", "message": str(e)}


@app.post("/file/edit")
async def edit_file(request: FileEditRequest):
    """Advanced file editing: overwrite, append, or patch (find & replace)."""
    try:
        global CURRENT_WORKSPACE
        if os.path.isabs(request.filepath):
            target_path = request.filepath
        else:
            target_path = os.path.join(CURRENT_WORKSPACE, request.filepath)

        if request.mode == "overwrite":
            parent_dir = os.path.dirname(target_path)
            if parent_dir and not os.path.exists(parent_dir):
                os.makedirs(parent_dir, exist_ok=True)
            with open(target_path, "w", encoding="utf-8") as f:
                f.write(request.content)
            return {"status": "SUCCESS", "path": target_path, "action": "overwritten"}

        elif request.mode == "append":
            with open(target_path, "a", encoding="utf-8") as f:
                f.write(request.content)
            return {"status": "SUCCESS", "path": target_path, "action": "appended"}

        elif request.mode == "patch":
            if not os.path.exists(target_path):
                return {"status": "ERROR", "message": f"File not found: {target_path}"}
            with open(target_path, "r", encoding="utf-8") as f:
                original = f.read()
            if request.patch_target and request.patch_target in original:
                new_content = original.replace(request.patch_target, request.patch_replacement or "", 1)
                with open(target_path, "w", encoding="utf-8") as f:
                    f.write(new_content)
                return {"status": "SUCCESS", "path": target_path, "action": "patched"}
            else:
                return {"status": "ERROR", "message": "Patch target not found in file."}
        else:
            return {"status": "ERROR", "message": f"Unknown mode: {request.mode}"}
    except Exception as e:
        return {"status": "ERROR", "message": str(e)}


@app.post("/file/read")
async def read_file(filepath: str = ""):
    """Read a file's content for AI context."""
    try:
        global CURRENT_WORKSPACE
        if os.path.isabs(filepath):
            target_path = filepath
        else:
            target_path = os.path.join(CURRENT_WORKSPACE, filepath)

        if not os.path.exists(target_path):
            return {"status": "ERROR", "message": "File not found"}
        with open(target_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return {"status": "SUCCESS", "content": content, "path": target_path}
    except Exception as e:
        return {"status": "ERROR", "message": str(e)}


# ─── TERMINAL ─────────────────────────────────────────────────────
@app.post("/terminal")
async def execute_command(request: TerminalRequest):
    try:
        global CURRENT_WORKSPACE
        cmd = request.command.strip()

        # Handle 'cd' commands to change CURRENT_WORKSPACE
        if cmd.lower().startswith("cd "):
            target = cmd[3:].strip().strip('"').strip("'")
            if os.path.isabs(target):
                new_dir = os.path.normpath(target)
            else:
                new_dir = os.path.normpath(os.path.join(CURRENT_WORKSPACE, target))
            if os.path.isdir(new_dir):
                CURRENT_WORKSPACE = new_dir
                settings = load_server_settings()
                settings["last_workspace"] = CURRENT_WORKSPACE
                save_server_settings(settings)
                return {"output": f"Changed directory to: {CURRENT_WORKSPACE}"}
            else:
                return {"output": f"ERROR: Directory not found: {new_dir}"}

        result = subprocess.run(
            cmd, shell=True, capture_output=True, text=True,
            timeout=30, cwd=CURRENT_WORKSPACE,
            encoding='utf-8', errors='replace'
        )
        output = result.stdout if result.stdout else result.stderr
        return {"output": output or "DONE (No output)"}
    except subprocess.TimeoutExpired:
        return {"output": "ERROR: Command timed out after 30 seconds."}
    except Exception as e:
        return {"output": f"CRITICAL_ERROR: {str(e)}"}


# ─── SYSTEM INSTRUCTION ──────────────────────────────────────────
def build_system_instruction(workspace_path):
    """Build the system instruction with workspace context."""
    workspace_files = []
    try:
        for root, dirs, filenames in os.walk(workspace_path):
            dirs[:] = [d for d in dirs if d not in {'.git', 'node_modules', 'dist', 'build', '__pycache__'}]
            for f in filenames:
                rel = os.path.relpath(os.path.join(root, f), workspace_path)
                workspace_files.append(rel)
            if len(workspace_files) > 100:
                break
    except:
        pass

    file_list_str = ", ".join(workspace_files[:50])
    context = retrieve_top_k("", k=4) if DOC_DB else "Рабочая область не проиндексирована."

    return (
        "Ты — VEY.AI, продвинутый русскоязычный ИИ-помощник для работы с файлами, кодом и текстом.\n"
        "ОТВЕЧАЙ МАКСИМАЛЬНО ПОДРОБНО, КРАСИВО (используя Markdown) И ТОЛЬКО НА РУССКОМ ЯЗЫКЕ.\n\n"
        "ОБЯЗАТЕЛЬНЫЕ ПРАВИЛА:\n"
        "1. Пиши связный текст. ЗАПРЕЩЕНО цитировать технические заголовки 'ФАЙЛ', '(ЧАСТЬ)', '[FILE_REQUEST]'.\n"
        "2. Для создания НОВОГО файла используй формат:\n"
        "   [FILE_REQUEST: путь/к/файлу.расширение]\n"
        "   ПОЛНОЕ СОДЕРЖИМОЕ ФАЙЛА\n"
        "   [/FILE_REQUEST]\n"
        "3. Для ИЗМЕНЕНИЯ существующего файла используй формат:\n"
        "   [FILE_EDIT: путь/к/файлу.расширение]\n"
        "   ПОЛНОЕ НОВОЕ СОДЕРЖИМОЕ ФАЙЛА (заменяет старое целиком)\n"
        "   [/FILE_EDIT]\n"
        "4. Для ЧАСТИЧНОГО ИЗМЕНЕНИЯ (патч) используй формат:\n"
        "   [FILE_PATCH: путь/к/файлу.расширение]\n"
        "   [FIND]\n"
        "   точный текст для замены\n"
        "   [/FIND]\n"
        "   [REPLACE]\n"
        "   новый текст замены\n"
        "   [/REPLACE]\n"
        "   [/FILE_PATCH]\n"
        "5. Ты МОЖЕШЬ создавать файлы ЛЮБОГО размера. Пиши полное содержимое без сокращений.\n"
        "6. Поддерживай русский язык в содержимом файлов.\n"
        "7. Если пользователь просит изменить файл — ОБЯЗАТЕЛЬНО используй один из тегов выше.\n\n"
        f"ТЕКУЩАЯ РАБОЧАЯ ПАПКА: {workspace_path}\n"
        f"ФАЙЛЫ В ПАПКЕ: {file_list_str}\n\n"
        f"--- LOCAL DATABANKS (Содержимое файлов) ---:\n{context}\n"
    )


# ─── STREAMING CHAT ──────────────────────────────────────────────
async def stream_groq_response(api_key, model, system_instruction, user_msg):
    """Stream response from Groq API with SSE events."""
    import json as json_mod

    yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Анализирую контекст рабочей области...'})}\n\n"
    await asyncio.sleep(0.1)

    yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Формирую запрос к нейросети...'})}\n\n"
    await asyncio.sleep(0.1)

    try:
        yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Получаю ответ от AI...'})}\n\n"

        res = requests.post("https://api.groq.com/openai/v1/chat/completions", headers={
            "Authorization": f"Bearer {api_key}"
        }, json={
            "model": model,
            "messages": [
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": user_msg}
            ],
            "max_tokens": 8192,
            "temperature": 0.3,
            "stream": True
        }, timeout=120, stream=True)

        if res.status_code != 200:
            error_body = res.text
            yield f"data: {json_mod.dumps({'type': 'error', 'content': f'GROQ Error ({res.status_code}): {error_body}'})}\n\n"
            yield f"data: {json_mod.dumps({'type': 'done'})}\n\n"
            return

        yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Генерирую ответ...'})}\n\n"

        full_content = ""
        for line in res.iter_lines():
            if not line:
                continue
            line_str = line.decode('utf-8', errors='replace')
            if line_str.startswith('data: '):
                data_str = line_str[6:]
                if data_str.strip() == '[DONE]':
                    break
                try:
                    chunk = json_mod.loads(data_str)
                    delta = chunk.get('choices', [{}])[0].get('delta', {})
                    content = delta.get('content', '')
                    if content:
                        full_content += content
                        yield f"data: {json_mod.dumps({'type': 'chunk', 'content': content})}\n\n"
                except:
                    pass

        yield f"data: {json_mod.dumps({'type': 'done', 'full_content': full_content})}\n\n"

    except requests.exceptions.Timeout:
        yield f"data: {json_mod.dumps({'type': 'error', 'content': 'Таймаут запроса к Groq API'})}\n\n"
        yield f"data: {json_mod.dumps({'type': 'done'})}\n\n"
    except Exception as e:
        yield f"data: {json_mod.dumps({'type': 'error', 'content': f'Ошибка: {str(e)}'})}\n\n"
        yield f"data: {json_mod.dumps({'type': 'done'})}\n\n"


async def stream_ollama_response(real_model, system_instruction, user_msg):
    """Stream response from Ollama with SSE events."""
    import json as json_mod

    yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Подключаюсь к Ollama...'})}\n\n"
    await asyncio.sleep(0.1)

    try:
        res = requests.post("http://localhost:11434/api/chat", json={
            "model": real_model,
            "messages": [
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": user_msg}
            ],
            "stream": True,
            "options": {"num_predict": 4096}
        }, timeout=180, stream=True)

        yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Генерирую ответ...'})}\n\n"

        full_content = ""
        for line in res.iter_lines():
            if not line:
                continue
            try:
                chunk = json_mod.loads(line)
                content = chunk.get("message", {}).get("content", "")
                if content:
                    full_content += content
                    yield f"data: {json_mod.dumps({'type': 'chunk', 'content': content})}\n\n"
                if chunk.get("done", False):
                    break
            except:
                pass

        yield f"data: {json_mod.dumps({'type': 'done', 'full_content': full_content})}\n\n"

    except Exception as e:
        yield f"data: {json_mod.dumps({'type': 'error', 'content': f'Ollama ошибка: {str(e)}'})}\n\n"
        yield f"data: {json_mod.dumps({'type': 'done'})}\n\n"


async def stream_local_response(system_instruction, user_msg):
    """Stream response from local LLM with SSE events."""
    import json as json_mod

    yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Загружаю локальную модель...'})}\n\n"
    await asyncio.sleep(0.1)

    try:
        model, tok = get_llm()
        import torch

        yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Токенизирую запрос...'})}\n\n"

        messages_for_llm = [
            {"role": "system", "content": system_instruction},
            {"role": "user", "content": user_msg}
        ]

        text = tok.apply_chat_template(messages_for_llm, tokenize=False, add_generation_prompt=True)
        inputs = tok([text], return_tensors="pt", truncation=True, max_length=4096).to(get_device())

        yield f"data: {json_mod.dumps({'type': 'thinking', 'content': 'Генерирую ответ...'})}\n\n"

        with torch.no_grad():
            outputs = model.generate(
                **inputs, max_new_tokens=2048,
                temperature=0.3, top_p=0.9, repetition_penalty=1.1
            )

        generated_ids = [output_ids[len(input_ids):] for input_ids, output_ids in zip(inputs.input_ids, outputs)]
        answer = tok.batch_decode(generated_ids, skip_special_tokens=True)[0].strip()

        # Send in chunks for visual effect
        chunk_size = 20
        for i in range(0, len(answer), chunk_size):
            chunk = answer[i:i+chunk_size]
            yield f"data: {json_mod.dumps({'type': 'chunk', 'content': chunk})}\n\n"
            await asyncio.sleep(0.01)

        yield f"data: {json_mod.dumps({'type': 'done', 'full_content': answer})}\n\n"

    except Exception as e:
        yield f"data: {json_mod.dumps({'type': 'error', 'content': f'Ошибка: {str(e)}'})}\n\n"
        yield f"data: {json_mod.dumps({'type': 'done'})}\n\n"


# ─── CHAT ENDPOINTS ──────────────────────────────────────────────
@app.post("/chat/stream")
async def chat_stream(request: ChatRequest):
    """Streaming chat endpoint - sends Server-Sent Events with thinking steps and response chunks."""
    last_msg = request.messages[-1]['content']

    # Build context
    try:
        build_vector_database(CURRENT_WORKSPACE)
    except:
        pass

    context = retrieve_top_k(last_msg, k=6)

    # Build workspace file list
    workspace_files = []
    try:
        for root, dirs, filenames in os.walk(CURRENT_WORKSPACE):
            dirs[:] = [d for d in dirs if d not in {'.git', 'node_modules', 'dist', 'build', '__pycache__'}]
            for f in filenames:
                rel = os.path.relpath(os.path.join(root, f), CURRENT_WORKSPACE)
                workspace_files.append(rel)
            if len(workspace_files) > 100:
                break
    except:
        pass

    file_list_str = ", ".join(workspace_files[:50])
    system_instruction = build_system_instruction(CURRENT_WORKSPACE)

    if request.model == "PYTHON_LOCAL_AI":
        return StreamingResponse(
            stream_local_response(system_instruction, last_msg),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            }
        )
    elif request.model.startswith("OLLAMA"):
        real_model = request.model.split(": ")[1] if ": " in request.model else "deepseek-coder"
        return StreamingResponse(
            stream_ollama_response(real_model, system_instruction, last_msg),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            }
        )
    else:
        # Groq
        if not request.api_key:
            async def error_stream():
                yield f"data: {json.dumps({'type': 'error', 'content': 'GROQ API Key отсутствует.'})}\n\n"
                yield f"data: {json.dumps({'type': 'done'})}\n\n"
            return StreamingResponse(
                error_stream(),
                media_type="text/event-stream",
            )
        return StreamingResponse(
            stream_groq_response(request.api_key, request.model, system_instruction, last_msg),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no",
            }
        )


@app.post("/chat")
async def chat(request: ChatRequest):
    """Standard (non-streaming) chat endpoint — kept for backwards compatibility."""
    try:
        last_msg = request.messages[-1]['content']

        build_vector_database(CURRENT_WORKSPACE)
        context = retrieve_top_k(last_msg, k=6)
        system_instruction = build_system_instruction(CURRENT_WORKSPACE)

        if request.model == "PYTHON_LOCAL_AI":
            model, tok = get_llm()
            import torch
            messages_for_llm = [
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": last_msg}
            ]

            text = tok.apply_chat_template(messages_for_llm, tokenize=False, add_generation_prompt=True)
            inputs = tok([text], return_tensors="pt", truncation=True, max_length=4096).to(get_device())

            with torch.no_grad():
                outputs = model.generate(
                    **inputs, max_new_tokens=2048,
                    temperature=0.3, top_p=0.9, repetition_penalty=1.1
                )

            generated_ids = [output_ids[len(input_ids):] for input_ids, output_ids in zip(inputs.input_ids, outputs)]
            answer = tok.batch_decode(generated_ids, skip_special_tokens=True)[0].strip()

            return {"answer": answer}

        elif request.model.startswith("OLLAMA"):
            real_model = request.model.split(": ")[1] if ": " in request.model else "deepseek-coder"
            try:
                res = requests.post("http://localhost:11434/api/chat", json={
                    "model": real_model,
                    "messages": [
                        {"role": "system", "content": system_instruction},
                        {"role": "user", "content": last_msg}
                    ],
                    "stream": False,
                    "options": {"num_predict": 4096}
                }, timeout=120).json()
                return {"answer": res.get("message", {}).get("content", "Error")}
            except:
                return {"answer": "OLLAMA CONNECTION FAILED."}
        else:
            # Groq
            if not request.api_key:
                return {"answer": "ERROR: GROQ API Key is missing."}
            try:
                res = requests.post("https://api.groq.com/openai/v1/chat/completions", headers={
                    "Authorization": f"Bearer {request.api_key}"
                }, json={
                    "model": request.model,
                    "messages": [
                        {"role": "system", "content": system_instruction},
                        {"role": "user", "content": last_msg}
                    ],
                    "max_tokens": 8192,
                    "temperature": 0.3
                }, timeout=60).json()
                if "error" in res:
                    return {"answer": f"GROQ Error: {res['error']['message']}"}
                return {"answer": res["choices"][0]["message"]["content"]}
            except Exception as e:
                return {"answer": f"GROQ SDK ERROR: {str(e)}"}
    except Exception as e:
        return {"answer": f"SYSTEM_ERROR: {str(e)}"}


# ─── METRICS ──────────────────────────────────────────────────────
@app.get("/metrics")
async def get_metrics():
    return {
        "cpu": psutil.cpu_percent(),
        "memory": psutil.virtual_memory().percent,
        "indexing": 100.0
    }

@app.get("/models/ollama")
async def list_ollama_models():
    try:
        response = requests.get("http://localhost:11434/api/tags", timeout=2)
        if response.status_code == 200:
            return {"models": [m['name'] for m in response.json().get('models', [])]}
        return {"models": []}
    except:
        return {"models": []}

@app.get("/workspace/files")
async def list_files(path: str = None):
    try:
        global CURRENT_WORKSPACE
        if path:
            CURRENT_WORKSPACE = path
            # Save last workspace
            settings = load_server_settings()
            settings["last_workspace"] = CURRENT_WORKSPACE
            save_server_settings(settings)

        root_dir = CURRENT_WORKSPACE
        files = []
        ignore_dirs = {'.git', 'node_modules', 'dist', 'build', '__pycache__', 'target', '.next', '.vey'}
        for root, dirs, filenames in os.walk(root_dir):
            dirs[:] = [d for d in dirs if d not in ignore_dirs]
            rel_path = os.path.relpath(root, root_dir)
            if rel_path != ".":
                files.append({"name": os.path.basename(root), "type": "directory", "path": rel_path})
            for f in filenames:
                file_path = f if rel_path == "." else os.path.join(rel_path, f)
                files.append({"name": f, "type": "file", "path": file_path})
        return {"files": files[:200]}
    except Exception as e:
        return {"error": str(e)}

@app.get("/workspace/current")
async def get_current_workspace():
    return {"path": CURRENT_WORKSPACE}

@app.post("/create-pdf")
async def create_pdf(data: dict):
    try:
        filename = data.get("filename", "document.pdf")
        content = data.get("content", "")
        title = data.get("title", "VEY.AI Document")

        from fpdf import FPDF
        import os

        class PDF(FPDF):
            def header(self):
                self.set_font('Arial', 'B', 15)
                self.cell(0, 10, title, 0, 1, 'C')
                self.ln(5)

            def footer(self):
                self.set_y(-15)
                self.set_font('Arial', 'I', 8)
                self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

        pdf = PDF()
        pdf.add_page()
        
        # Try to find a font that supports Cyrillic
        font_paths = [
            "C:\\Windows\\Fonts\\arial.ttf",
            "C:\\Windows\\Fonts\\tahoma.ttf",
            "C:\\Windows\\Fonts\\times.ttf"
        ]
        
        font_loaded = False
        for fp in font_paths:
            if os.path.exists(fp):
                try:
                    pdf.add_font("CustomFont", "", fp)
                    pdf.set_font("CustomFont", size=12)
                    font_loaded = True
                    break
                except:
                    continue
        
        if not font_loaded:
            pdf.set_font("Arial", size=12)

        # Process multi-line content
        for line in content.split('\n'):
            pdf.multi_cell(0, 10, line)

        output_path = os.path.join(CURRENT_WORKSPACE, filename)
        pdf.output(output_path)
        
        return {"status": "ok", "path": output_path}
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.post("/terminal")
async def run_terminal(data: dict):
    try:
        command = data.get("command", "")
        if not command:
            return {"output": "No command provided"}
            
        import subprocess
        result = subprocess.run(command, shell=True, capture_output=True, text=True, cwd=CURRENT_WORKSPACE)
        output = result.stdout if result.stdout else result.stderr
        return {"output": output}
    except Exception as e:
        return {"output": f"Error: {str(e)}"}


if __name__ == "__main__":
    print("Vey AI Core starting on http://127.0.0.1:8000")
    uvicorn.run(app, host="127.0.0.1", port=8000)
